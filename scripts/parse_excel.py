#!/usr/bin/env python3
"""
Lê o arquivo GESTÃO_TOLVA (.xlsm) e gera os JSONs usados pelo site.

Lógica replicada da planilha (aba GERAL, Tabela6):
  ULTIMA   = MAX(coluna "2025", coluna "OS")   -> data da última troca
  PRÓXIMA  = ULTIMA + FREQU.                   -> calculado no navegador (JS)
  DIAS RESTANTES = PRÓXIMA - HOJE              -> calculado no navegador (JS)
  STATUS   = <2 dias: PENDENTE | >=20 dias: DENTRO DO PRAZO | senão: ATENÇÃO

Como PRÓXIMA/DIAS RESTANTES/STATUS dependem da data de hoje, eles NÃO são
gravados no JSON — o site recalcula isso sozinho a cada carregamento,
usando a data do dia. Só precisamos reprocessar o Excel quando uma troca
de filtro realmente acontece (a data "ULTIMA" muda).
"""

import json
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

EXCEL_EPOCH = datetime(1899, 12, 30)  # mesma referência de data serial do Excel

REPO_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = REPO_ROOT / "data"


def to_serial(value):
    """Converte datetime ou número em um 'serial' comparável (estilo Excel)."""
    if isinstance(value, datetime):
        return (value - EXCEL_EPOCH).days
    if isinstance(value, (int, float)):
        return value
    return None


def serial_to_iso(serial):
    return (EXCEL_EPOCH + timedelta(days=int(serial))).strftime("%Y-%m-%d")


def find_excel_file():
    uploads_dir = REPO_ROOT / "uploads"
    candidates = sorted(uploads_dir.glob("*.xlsm")) + sorted(uploads_dir.glob("*.xlsx"))
    if not candidates:
        sys.exit("Nenhum arquivo .xlsm/.xlsx encontrado em uploads/")
    # pega o mais recentemente modificado
    return max(candidates, key=lambda p: p.stat().st_mtime)


def parse_geral(ws):
    """Lê a tabela Tabela6 (B2:K23) da aba GERAL."""
    itens = []
    row = 3
    while True:
        setor = ws.cell(row=row, column=2).value  # B
        if setor is None or str(setor).strip() == "":
            break

        tag = ws.cell(row=row, column=3).value        # C
        filtro = ws.cell(row=row, column=4).value      # D
        frequencia = ws.cell(row=row, column=5).value  # E
        col_2025 = ws.cell(row=row, column=6).value    # F
        col_os = ws.cell(row=row, column=7).value      # G

        serial_2025 = to_serial(col_2025)
        serial_os = to_serial(col_os)
        serials = [s for s in (serial_2025, serial_os) if s is not None]

        ultima_iso = serial_to_iso(max(serials)) if serials else None

        itens.append({
            "setor": str(setor).strip(),
            "tag": tag,
            "filtro": filtro,
            "frequencia_dias": frequencia,
            "ultima_troca": ultima_iso,
        })
        row += 1

    return itens


def parse_estoque(ws):
    """Lê a aba Plan4 (estoque/medidas de filtros), se existir dado."""
    estoque = []
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=3, max_col=5):
        quant, medida, descricao = (c.value for c in row)
        if quant is None and medida is None and descricao is None:
            continue
        estoque.append({
            "quantidade": quant,
            "medida": medida,
            "descricao": descricao,
        })
    return estoque


def main():
    excel_path = find_excel_file()
    print(f"Lendo: {excel_path.name}")

    wb = openpyxl.load_workbook(excel_path, data_only=True, keep_vba=True)

    itens = parse_geral(wb["GERAL"])
    estoque = parse_estoque(wb["Plan4"]) if "Plan4" in wb.sheetnames else []

    today = datetime.now().strftime("%Y-%m-%d")

    payload = {
        "gerado_em": today,
        "arquivo_origem": excel_path.name,
        "itens": itens,
        "estoque_filtros": estoque,
    }

    DATA_DIR.mkdir(exist_ok=True)
    (DATA_DIR / "latest.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (DATA_DIR / f"{today}.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"OK: {len(itens)} itens e {len(estoque)} linhas de estoque gravados em data/latest.json")


if __name__ == "__main__":
    main()
